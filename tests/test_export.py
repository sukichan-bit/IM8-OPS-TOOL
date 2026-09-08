import sys

BASE = r"C:/Users/suki.chan_prenetics/Desktop/Claude/im8-open-SO-fulfillment-checking-tools"
sys.path.insert(0, BASE)

import pandas as pd
import openpyxl

from lib.io_utils import to_excel_bytes
from lib.task_a import compute_production_requirement

onhand_path = f"{BASE}/samples/20260723 H007 on-hand (as of 1705).xlsx"
soso_path = f"{BASE}/samples/20260723 H007 Open SO (Jul22).xlsx"

onhand_df = pd.read_excel(onhand_path)
onhand_df.columns = [str(c).strip() for c in onhand_df.columns]

so_df = pd.read_excel(soso_path, sheet_name="SO Status", header=3)
so_df.columns = [str(c).strip() for c in so_df.columns]
requested_df = so_df[so_df["Remarks"] == "IT - to rerun fulfillment"].copy()

requested_cols = {"item": "Item number", "warehouse": "Warehouse", "qty": "Quantity"}
onhand_cols = {"item": "Item number", "warehouse": "Warehouse", "available": "Available physical", "product_name": "Product name"}

per_wh, diag = compute_production_requirement(requested_df, onhand_df, requested_cols, onhand_cols, show_all_rows=True)

row_colors = {
    wh: ["" if r["Item number"] == "TOTAL" else ("red" if r["To produce"] > 0 else "green") for _, r in table.iterrows()]
    for wh, table in per_wh.items()
}
xlsx_bytes = to_excel_bytes(per_wh, row_colors=row_colors)

out_path = f"{BASE}/tests/_sample_output_task_a.xlsx"
with open(out_path, "wb") as f:
    f.write(xlsx_bytes)

# Read it back and verify structure + red/green highlighting + sort order.
wb = openpyxl.load_workbook(out_path)
print("Sheets in output:", wb.sheetnames)
assert set(wb.sheetnames) == set(per_wh.keys())

ws = wb["OPS-WH03"]
red_count, green_count = 0, 0
last_to_produce = None
to_produce_col = [c[0].value for c in ws.iter_cols(min_row=1, max_row=1)].index("To produce") + 1
for i, row in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row - 1)):  # exclude TOTAL row
    fill = row[0].fill
    rgb = fill.fgColor.rgb if fill and fill.fgColor else None
    if rgb == "FFFFC7CE":
        red_count += 1
    elif rgb == "FFC6EFCE":
        green_count += 1
    to_produce = row[to_produce_col - 1].value
    if last_to_produce is not None:
        assert to_produce <= last_to_produce, f"expected descending To produce sort, row {i}: {to_produce} > {last_to_produce}"
    last_to_produce = to_produce

totals_row = list(ws.iter_rows(min_row=ws.max_row, max_row=ws.max_row))[0]
assert totals_row[0].fill.fgColor is None or totals_row[0].fill.fgColor.rgb not in ("FFFFC7CE", "FFC6EFCE"), "TOTAL row should not be colored"

print(f"OPS-WH03: {red_count} red rows, {green_count} green rows")
assert red_count > 0, "expected some red (needs production) rows"
assert green_count > 0, "expected some green (no shortfall) rows"

print("EXPORT TEST PASSED")
